from docx import Document
doc = Document(r'Files/Conference template.docx')
print('=== PARAGRAPH STYLES IN USE ===')
for i, para in enumerate(doc.paragraphs):
    print(f'[{i:03d}] style={para.style.name!r:35s} | {para.text[:100]}')
print()
print('=== TABLES ===')
for t_idx, table in enumerate(doc.tables):
    print(f'Table {t_idx}: {len(table.rows)} rows x {len(table.columns)} cols')
    for r in table.rows:
        for c in r.cells:
            print(f'  | {c.text[:60]}', end='')
        print()
